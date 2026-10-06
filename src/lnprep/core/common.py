"""
Shared helpers, 3-zone architecture rules, and format specification loaders.

Centralises PPTX, LECTURE_NOTES_GUIDE.md, and NOTES_FORMAT.md handling.
Also provides prompt-side rule text, canonical gold examples, and format spec renderers.
"""

from __future__ import annotations

import os
import re
from typing import Any, Dict, List, Optional, Tuple

import yaml
from pptx.enum.shapes import MSO_SHAPE_TYPE

# === 3-ZONE CONSTANTS & STRUCTURAL DELIMITERS ===

ZONE_SPEAKER = "--- SPEAKER NOTES ---"
ZONE_VISUAL = "--- VISUAL DECONSTRUCTION ---"
ZONE_LECTURE = "--- LECTURE NOTES ---"

SLIDE_START = "## Slide Content\n\n"
SLIDE_END = "\n---\n\n## Required Output Format"
BATCH_SLIDE_SEP = "\n\n---\n\n"

FORMAT_FILENAME = "NOTES_FORMAT.md"
GUIDE_FILENAME = "LECTURE_NOTES_GUIDE.md"


# === GUIDE HELPERS ===

def find_guide_file(pptx_path: str) -> Optional[str]:
    """Walk up from the PPTX path to find LECTURE_NOTES_GUIDE.md.

    Searches up to 5 parent directories starting from the PPTX's own folder.
    Returns the absolute path to the guide, or None if none is found.
    """
    current = os.path.dirname(os.path.abspath(pptx_path))
    for _ in range(5):
        guide_path = os.path.join(current, GUIDE_FILENAME)
        if os.path.exists(guide_path):
            return guide_path
        parent = os.path.dirname(current)
        if parent == current:
            break
        current = parent
    return None


def parse_guide(guide_path: Optional[str]) -> Dict[str, Any]:
    """Parse LECTURE_NOTES_GUIDE.md for module context.

    Extracts every yaml fenced block's key: value pairs into a flat dict,
    plus _session_topics_raw and _has_notes_structure.
    Bare fences are read with a strict lowercase snake_case filter.
    Returns {} if the guide is missing or unreadable.
    """
    if not guide_path or not os.path.exists(guide_path):
        return {}

    try:
        with open(guide_path, "r", encoding="utf-8") as f:
            content = f.read()
    except OSError:
        return {}

    context: Dict[str, Any] = {}

    def _harvest(block: str, strict: bool) -> None:
        for line in block.strip().split("\n"):
            line = line.strip()
            if ":" not in line or line.startswith("#"):
                continue
            key, _, val = line.partition(":")
            key = key.strip().strip('"')
            if strict and not re.match(r"^[a-z][a-z0-9_]{1,40}$", key):
                continue
            val = val.strip().strip('"')
            if val:
                context[key] = val

    for block in re.findall(r"```yaml\n(.*?)```", content, re.DOTALL):
        _harvest(block, strict=False)
    for block in re.findall(r"```[ \t]*\n(.*?)```", content, re.DOTALL):
        _harvest(block, strict=True)

    session_match = re.search(r"## Session Topics.*?```yaml\n(.*?)```", content, re.DOTALL)
    if session_match:
        context["_session_topics_raw"] = session_match.group(1)

    if "## Notes Section Structure" in content:
        context["_has_notes_structure"] = True

    return context


# === SLIDE HELPERS ===

def get_slide_title(slide: Any) -> str:
    """Extract the slide title, falling back to the first text shape.

    Returns '(No title)' when the slide has no usable title or text.
    """
    if getattr(slide.shapes, "title", None) and slide.shapes.title.text.strip():
        return slide.shapes.title.text.strip()
    for shape in slide.shapes:
        if getattr(shape, "has_text_frame", False) and shape.text_frame.text.strip():
            return shape.text_frame.text.strip()[:80]
    return "(No title)"


def extract_slide_body_text(slide: Any) -> str:
    """Extract all body text from a slide, preserving bullet indentation.

    The title shape is skipped. Recursively traverses GROUP shapes.
    """
    lines: List[str] = []

    def _walk(shapes: Any) -> None:
        for shape in shapes:
            if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                try:
                    _walk(shape.shapes)
                except (AttributeError, ValueError):
                    pass
                continue
            if not getattr(shape, "has_text_frame", False):
                continue
            if shape == getattr(slide.shapes, "title", None):
                continue
            for para in shape.text_frame.paragraphs:
                text = para.text.strip()
                if text:
                    level = getattr(para, "level", 0)
                    prefix = "  " * level + "• "
                    lines.append(f"{prefix}{text}")

    _walk(slide.shapes)
    return "\n".join(lines)


def has_visual_content(slide: Any) -> bool:
    """Check if a slide has tables, charts, images, groups, or SmartArt."""
    for shape in slide.shapes:
        if getattr(shape, "has_table", False):
            return True
        if shape.shape_type == MSO_SHAPE_TYPE.CHART:
            return True
        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
            return True
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            return True
    try:
        for rel in slide.part.rels.values():
            if "diagram" in rel.reltype:
                return True
    except Exception:
        pass
    return False


def get_slide_notes_text(slide: Any) -> str:
    """Safely extract slide notes text.

    Guards against slides where has_notes_slide is True but notes_text_frame
    is None or raises AttributeError/ValueError.
    """
    if not getattr(slide, "has_notes_slide", False):
        return ""
    try:
        tf = slide.notes_slide.notes_text_frame
    except (AttributeError, ValueError):
        return ""
    if tf is None:
        return ""
    return tf.text.strip()


# === RULES & SPEC SCHEMA ===

RETIRED_GUIDE_KEYS = {
    "notes_format": "superseded by NOTES_FORMAT.md (the file, not a key)",
    "sbc_depth": "superseded by sbc_quality.pass_threshold + sbc_fields",
    "sbc_research_enrich": "superseded by sbc_fields + vocabulary",
    "research_enrichment": "no consumer — Feature 7 is agent-driven, not scripted",
    "include_videos": "no consumer — Feature 8 is agent-driven, not scripted",
    "zone_a_script": "superseded by zones.speaker_notes",
    "zone_b_cheatsheet": "superseded by zones.lecture_notes",
    "zone_c_reference": "superseded by zones.sbc",
}

FORMAT_SPEC_SCHEMA = {
    "format_version": "live",
    "status": "live",
    "module_code": "reserved",
    "delivery_mode": "live",
    "zones": "live",
    "sbc_header": "live",
    "zone_a_labels": "live",
    "zone_a_hook_archetypes": "live",
    "require_teacher_guidance": "live",
    "sbc_fields": "live",
    "sbc_quality": "live",
    "vocabulary": "live",
    "audit": "live",
    "fonts": "live",
    "engagement": "live",
}

PARAGRAPH_RULES = {
    "key_point": {"min": 1, "max": 3, "prose_only": True},
    "plain_english": {"min": 1, "max": 3, "prose_only": True},
    "default": {"min": 2},
}

KEY_POINT_RULE = (
    "KEY POINT: 1–3 paragraphs of continuous prose, separated by blank lines. No bullets or lists. "
    "Para 1 (required): the core principle — what it is and why it matters. "
    "Para 2 (if needed): the mechanism, trade-off, or organisational friction. "
    "Para 3 (if needed): the strategic consequence and assessment relevance. "
    "Stop when the idea is complete; never pad to reach 3."
)

PLAIN_ENGLISH_RULE = (
    "Plain English: 1–3 paragraphs of jargon-free prose, separated by blank lines. No bullets. "
    "Para 1 (required): the definition and how it works, in everyday words. "
    "Para 2 (required unless the metaphor fits in Para 1): an explicit pedagogical metaphor "
    '("Think of it like..."). '
    "Para 3 (optional): where it breaks down — a boundary condition or edge case."
)

DEEP_RESEARCH_EVIDENCE_RULE = (
    "Put a DOI (doi:10.xxxx/...) after every journal citation, and an ISBN or publisher URL after "
    "every book, so verify_references.py can check it. NEVER fabricate or cite from memory: an "
    "unverifiable citation blocks write-back."
)

EXAMPLE_SOURCES_RULE = (
    "Every figure or company fact carries an inline (Organisation/Author, Year) marker, and the "
    'field ENDS with one line: Sources: <Org> (<Year>) "<Title>", <URL or DOI>; <next source>. '
    "State the data year of every figure (e.g. FY2024). Prefer primary sources — annual reports, "
    "investor relations, regulators (MPA, SingStat, MAS), peer-reviewed papers. No Wikipedia, AI "
    "summaries, undated blogs, or search-engine redirect links. A Concrete Example without a "
    "Sources line blocks write-back."
)

VISUAL_DECONSTRUCTION_RULE = (
    "VISUAL DECONSTRUCTION: Required whenever the slide contains a diagram, framework, chart, "
    "table, SmartArt, process flow, or image; omit entirely for text-only/agenda slides. "
    "Structure as an explicit top-level zone between --- SPEAKER NOTES --- and --- LECTURE NOTES ---. "
    "Include four micro-elements:\n"
    "1. 🎯 VISUAL OBJECTIVE: One sentence stating what core relationship, mechanism, or data insight this graphic proves.\n"
    "2. 👁️ GAZE DIRECTION (Where to point first): Physical pointer/eye guidance on where students must look first (axes, inputs, baseline quadrant).\n"
    "3. 🔄 STEP-BY-STEP WALK-THROUGH: Numbered steps (Step 1 Baseline/Inputs -> Step 2 Mechanism/Flow -> Step 3 Friction/Inflection Point -> Step 4 Punchline Synthesis).\n"
    "4. 💡 VISUAL PUNCHLINE: One-sentence managerial or theoretical takeaway revealed directly by the visual."
)

GOLD_KEY_POINT_EXAMPLE = """\
KEY POINT: "A stakeholder is anyone who can affect, or is affected by, what the firm does, and managers rarely have the time or resources to serve all of them equally.

That is why salience matters. Groups with power, legitimacy and urgency get attention first, and a group can move up the list quickly when it gains one of those attributes.

For your Task 2 analysis, the strongest answers rank stakeholders and explain how that ranking changes the firm's operational priorities.\""""

GOLD_VISUAL_DECONSTRUCTION_EXAMPLE = """\
--- VISUAL DECONSTRUCTION ---
🎯 VISUAL OBJECTIVE:
Demonstrates how small stochastic demand variances amplify exponentially as information moves upstream through supply chain tiers.

👁️ GAZE DIRECTION (Where to point first):
"Direct students' eyes to the far left of the diagram: Consumer Retail Sales (the flat, mild wave)."

🔄 STEP-BY-STEP WALK-THROUGH:
- Step 1 [Consumer to Retailer]: "Point to the retail order line: notice the mild ±5% fluctuation. Explain that batch ordering and lead-time safety buffers slightly exaggerate this wave."
- Step 2 [Retailer to Distributor to Manufacturer]: "Trace your finger to the middle tiers: show how the amplitude doubles at the distributor level and triples at the manufacturing tier due to order batching and price speculation."
- Step 3 [The Tier-2 / Raw Material Whiplash]: "Point to the violent, chaotic wave at the far right (Raw Materials). Highlight that component suppliers experience wild boom-and-bust cycles despite stable end-consumer consumption."
- Step 4 [Punchline Synthesis]: "Synthesize the visual rule: 'The distortion is not caused by unpredictable consumers; it is caused by internal information lag and independent buffer hoarding across silos.'"

💡 VISUAL PUNCHLINE:
"The distortion is not caused by unpredictable consumers; it is caused by internal information lag and independent buffer hoarding across silos.\""""

GOLD_SBC_EXAMPLE = """\
• SLIDE BODY COVERAGE — MAJOR CORPORATE STAKEHOLDERS:

Major Corporate Stakeholders:

Plain English: Major corporate stakeholders are any group, individual, or ecological entity that is affected by, or can influence, what an organisation does, how it transforms resources, and what it produces.

Think of it like building a major airport: you cannot satisfy only the airline passengers; you must also answer to air traffic controllers, the neighbourhoods living with the noise, the municipal drainage system, and the retail staff in the terminal.

The idea breaks down when every group is treated as equally important. A firm that tries to satisfy everyone equally ends up prioritising no one, which is why stakeholder salience (who actually counts, and when) matters.

Deep Research:
- Classical Foundation: Freeman (1984) *Strategic Management: A Stakeholder Approach* founded stakeholder theory (ISBN 9780273019138); Mitchell, Agle & Wood (1997) in *Academy of Management Review* explained which stakeholders managers actually attend to through salience: power, legitimacy and urgency (doi:10.5465/amr.1997.9711022105).
- Contemporary Frontier (2024-2026): Matthews et al. (2025) in *Journal of Management* review AI, algorithms and robots through a stakeholder lens, showing that intelligent machines often create both advantages and disadvantages for different stakeholders (doi:10.1177/01492063241311855).

Concrete Example: For example, PSA's flagship Singapore terminal handled a record 40.9 million TEUs in 2024, up 5.5% (PSA International, 2025), volume that depends on keeping shipping lines, port unions, regulators and coastal communities aligned at the same time.

Globally, Nike, long the subject of sweatshop allegations, in 2005 became the first major footwear and apparel company to publish the names and locations of its more than 700 active contract factories (Nike Inc., 2005; Teather, 2005).

Sources: PSA International (2025) "PSA International's 2024 Container Throughput Performance", https://www.singaporepsa.com/2025/01/16/psa-internationals-2024-container-throughput-performance/; Nike Inc. (2005) "Nike Publishes List of Global Contract Factories in Push for Greater Transparency", https://csrwire.com/press-release/nike-publishes-list-of-global-contract-factories-in-push-for-greater-transparency-and-collaboration-to-improve-footwear-and-apparel-industry-labor-conditions/; Teather, D. (2005) "Nike lists abuses at Asian factories", The Guardian, https://www.theguardian.com/business/2005/apr/14/ethicalbusiness.money

Bigger Picture: Strategically, stakeholder legitimacy works like a licence to operate: once a firm loses it, regulators, customers and investors can withdraw support at the same time, so the damage compounds.

The trade-off is attention. Managers cannot serve every group fully, so the advantage goes to firms that read salience early and act before a latent stakeholder becomes a definitive one.

Assessment Link: Use stakeholder salience analysis in Task 2 to justify why your selected MNC must address specific operational sustainability challenges.

Distinction-level answers classify each stakeholder by power, legitimacy and urgency and show how that ranking changes the firm's priorities; listing stakeholders without ranking them stays descriptive.

Manager's So What: Actionable takeaway: build a stakeholder salience map that tracks latent, expectant and definitive stakeholders across every sourcing region.

Review the map whenever a stakeholder gains a new attribute, such as a community group gaining media attention and with it urgency, because that is the moment it moves up the priority list."""

DEFAULT_FORMAT_SPEC: Dict[str, Any] = {
    "format_version": 1,
    "status": "active",
    "module_code": None,
    "delivery_mode": None,
    "zones": {
        "speaker_notes": True,
        "visual_deconstruction": True,
        "lecture_notes": True,
        "sbc": True,
        "speaker_first": True,
    },
    "sbc_header": "SLIDE BODY COVERAGE",
    "zone_a_labels": [
        "BRIDGE",
        "NARRATIVE OPENING / HOOK",
        "KEY POINT",
        "SOCRATIC PROMPT / COLD CALL",
        "COMMON STUDENT MISCONCEPTION",
        "ACTIVITY",
        "TEACHER GUIDANCE & EXPECTED ANSWERS",
        "VIDEO CUE",
        "SCRIPTED EXPLANATION",
        "TRANSITION",
    ],
    "zone_a_hook_archetypes": ["crisis", "paradox", "trade_off", "metric"],
    "require_teacher_guidance": True,
    "sbc_fields": [
        {
            "label": "Plain English",
            "require_metaphor": True,
            "lead_in": "Think of it like",
            "prompt": PLAIN_ENGLISH_RULE,
        },
        {
            "label": "Deep Research",
            "require_dual_anchor": True,
            "prompt": (
                "Deep Research: Dual-anchor academic foundation:\n"
                "- Classical Foundation: Seminal citation (Author, Year) establishing the baseline operational mechanism.\n"
                "- Contemporary Frontier (2024-2026): Recent peer-reviewed research (Author, 2024-2026) updating the concept for digital, geopolitical, or resilient supply network contexts.\n"
                + DEEP_RESEARCH_EVIDENCE_RULE
            ),
        },
        {
            "label": "Concrete Example",
            "lead_in": "For example, ",
            "require_dual_geography": True,
            "prompt": (
                'Concrete Example: Dual-geography real-world illustration. MUST start with "For example, "\n'
                "and pair an APAC / Singapore local ecosystem anchor (e.g., PSA Tuas Mega Port, SIA Cargo,\n"
                "Grab, Keppel, TSMC, BYD, Shein, Dyson Singapore) alongside a global benchmark (Apple,\n"
                "Amazon, Toyota, FedEx), complete with specific quantifiable metrics (volume, cost, %, time).\n"
                + EXAMPLE_SOURCES_RULE
            ),
        },
        {
            "label": "Bigger Picture",
            "lead_in": "Strategically, ",
            "prompt": (
                'Bigger Picture: MUST start with "Strategically, " and articulate wider supply chain\n'
                "trade-offs, systemic ripple effects, and competitive advantage implications."
            ),
        },
        {
            "label": "Assessment Link",
            "prompt": (
                "Assessment Link: Explicitly connects the concept to module coursework brief sections,\n"
                "task word counts, evaluation criteria, or case report questions."
            ),
        },
        {
            "label": "Manager's So What",
            "lead_in": "Actionable takeaway: ",
            "prompt": (
                'Manager\'s So What: MUST start with "Actionable takeaway: " providing a pragmatic executive\n'
                "rule of thumb or decision protocol."
            ),
        },
    ],
    "sbc_quality": {
        "pass_threshold": 4,
        "markers": [
            ("definition", r"(?i)(?:is|means|refers to|defined as|describes)"),
            ("research_context", r"(?i)(?:\b(?:19|20)\d{2}\b|et al\.?|journal|paper|study|research)"),
            ("concrete_example", r"(?i)(?:for example|think of|imagine|consider|such as|e\.g\.|like the)"),
            ("bigger_picture", r"(?i)(?:strategically|why.*matter|consequence|implication|if.*wrong|critical)"),
            ("assessment_link", r"(?i)(?:assignment|exam|report|question|apply|calculate|analyse)"),
            ("so_what", r"(?i)(?:so what|manager.*should|leader.*do|actionable|takeaway|bottom line)"),
        ],
    },
    "vocabulary": {
        "regional_anchors": [
            "PSA Tuas Mega Port",
            "SIA Cargo",
            "Grab",
            "Keppel",
            "TSMC",
            "BYD",
            "Shein",
            "Dyson Singapore",
        ],
        "global_benchmarks": ["Apple", "Amazon", "Toyota", "FedEx"],
        "stopwords_extra": [],
    },
    "audit": {
        "stale_threshold_years": 5,
        "require_all_fields": False,
        "verify_formatting": False,
        "paragraphs": PARAGRAPH_RULES,
    },
    "fonts": {"timing": 1100, "header": 1000, "body": 1000, "sub": 900},
    "engagement": {
        "max_gap_minutes": 12,
        "required_formats": ["chat", "breakout_deliverable"],
    },
    "_prompt": {
        "sbc_intro": (
            "### SBC Quality Bar (lecture-notes-prep v2.12.0 — MO9529 Gold Standard)\n\n"
            "Every item in a Slide Body Coverage block must be a STANDALONE BLOCK headed by\n"
            "the exact slide bullet text, followed by ALL SIX canonical labeled fields separated\n"
            "by blank lines (Plain English: 1–3 paragraphs of prose; the other five fields: 2–3 paragraphs where feasible):\n\n"
            "<Exact Slide Bullet / Item Header>:"
        ),
        "sbc_outro": (
            "Header Matching Rule:\n"
            "The item header MUST match the slide bullet or shape text exactly to ensure 100%\n"
            "coverage in automated SBC audit passes.\n"
        ),
        "zone_a_intro": (
            "### Zone A High-Engagement Narrative Hook Spec (lecture-notes-prep v2.12.0 — MO9529 Gold Standard)\n\n"
            "Every slide must begin with the timing and pedagogical focus line:\n"
            "⏱ X.X min | 🎯 <Hook Title / Pedagogical Focus>\n\n"
            "Followed by an engaging, verbatim narrative opening hook:\n"
            'NARRATIVE OPENING / HOOK: "..."\n\n'
            "NARRATIVE OPENING / HOOK can span 1–2 paragraphs.\n\n"
            + KEY_POINT_RULE
            + "\n\nSBC fields: Plain English 1–3 paragraphs of prose; the other five fields 2–3 paragraphs where feasible, to unpack operational mechanisms, empirical regional/global data, and executive trade-offs.\n\n"
            "TEACHER GUIDANCE & EXPECTED ANSWERS: must be followed by a blank line and 5-space indented sub-bullets with dashes (     - Expected:,      - Guidance:,      - Assessment Link:,      - Key Insight:). In PPTX notes, these render as clean Level-1 indented bullets with bold labels.\n\n"
            "The hook MUST use one of the 4 Executive Pedagogical Hook Archetypes:"
        ),
        "zone_a_archetypes": [
            (
                "1. High-Stakes Operational Crisis / 3 AM Dilemma:\n"
                "   Place the student into the shoes of a supply chain director facing an immediate breakdown\n"
                '   (e.g., "Imagine it\'s 3:00 AM on a Friday, and a cyberattack paralyzes PSA Tuas\'s automated cranes...").'
            ),
            (
                "2. Counter-Intuitive Operational Paradox:\n"
                "   Highlight a real-world business decision that seems irrational on the surface\n"
                '   (e.g., "Why did Zara deliberately pay 300% more in air freight to fly dresses from Spain to Singapore, and end up doubling profit margins?").'
            ),
            (
                "3. Provocative Trade-Off / 'Choose Your Poison':\n"
                "   Force students to confront harsh operational realities and conflicting constraints\n"
                '   (e.g., "If your CEO demands a 40% inventory cut while guaranteeing 99% next-day fulfillment, what fundamental law are they asking you to break?").'
            ),
            (
                "4. Startling Reality-Check / Eye-Opening Metric:\n"
                "   Drop a verifiable industry statistic that shatters naive assumptions\n"
                '   (e.g., "72% of chief supply chain officers admit zero visibility past Tier-1. Today, we discover why a $2 microchip shut down global car manufacturing for months.").'
            ),
        ],
        "zone_a_outro": 'CRITICAL: Never use dry, flat openings like "On this slide, we will discuss...". Make it dramatic, provocative, and delivery-ready.\n',
    },
}


def find_format_files(pptx_path: str, max_levels: int = 5) -> List[str]:
    """Collect every NOTES_FORMAT.md from the deck's folder up to module root."""
    found: List[str] = []
    current = os.path.dirname(os.path.abspath(pptx_path))
    for _ in range(max_levels):
        candidate = os.path.join(current, FORMAT_FILENAME)
        if os.path.exists(candidate):
            found.append(candidate)
        parent = os.path.dirname(current)
        if parent == current:
            break
        current = parent
    return found


def _deep_merge(base: Dict[str, Any], overlay: Dict[str, Any]) -> Dict[str, Any]:
    """Recursively merge overlay onto base; overlay wins per leaf key."""
    out = dict(base)
    for key, val in overlay.items():
        if isinstance(val, dict) and isinstance(out.get(key), dict):
            out[key] = _deep_merge(out[key], val)
        else:
            out[key] = val
    return out


def spec_dead_knobs(spec: Dict[str, Any]) -> List[Tuple[str, str]]:
    """Report spec keys that no consumer reads."""
    findings: List[Tuple[str, str]] = []
    for key, val in spec.items():
        if key == "_prompt":
            continue
        if key in RETIRED_GUIDE_KEYS:
            findings.append((key, f"retired — {RETIRED_GUIDE_KEYS[key]}"))
            continue
        status = FORMAT_SPEC_SCHEMA.get(key)
        if status is None:
            findings.append((key, "unknown — not in FORMAT_SPEC_SCHEMA, so nothing reads it"))
        elif status == "reserved" and val is not None:
            findings.append((key, "reserved — declared but no consumer yet"))
    return findings


def _parse_spec_text(raw: str) -> Dict[str, Any]:
    """Parse a NOTES_FORMAT.md into a dict."""
    fence = re.search(r"```ya?ml[ \t]*\n(.*?)```", raw, re.DOTALL)
    return yaml.safe_load(fence.group(1) if fence else raw) or {}


def load_format_spec(pptx_path: str) -> Tuple[Dict[str, Any], Dict[str, Any]]:
    """Resolve the format spec for a deck.
    Walks up from the deck collecting NOTES_FORMAT.md files and merges nearest-wins.
    Returns (spec, report).
    """
    paths = find_format_files(pptx_path)
    report: Dict[str, Any] = {
        "found": list(paths),
        "bound": [],
        "draft": [],
        "errors": [],
        "warnings": [],
        "sources": {},
    }

    if not paths:
        return dict(DEFAULT_FORMAT_SPEC), report

    loaded_by_path: Dict[str, Any] = {}
    for path in paths:
        try:
            with open(path, "r", encoding="utf-8") as f:
                raw = f.read()
            loaded = _parse_spec_text(raw)
        except Exception as exc:
            report["errors"].append(f"{path}: {exc}")
            continue

        if not isinstance(loaded, dict):
            report["errors"].append(f"{path}: not a YAML mapping")
            continue

        status = str(loaded.get("status", "active")).strip().lower()
        if status == "draft":
            report["draft"].append(path)
            continue

        loaded_by_path[path] = loaded
        report["bound"].append(path)

    merged: Dict[str, Any] = {}
    for path in report["bound"]:
        for key, val in loaded_by_path[path].items():
            if key in merged and key not in ("sbc_fields", "zone_a_labels"):
                continue
            merged[key] = val
            report["sources"][key] = path

    spec = _deep_merge(DEFAULT_FORMAT_SPEC, merged)
    for key in ("sbc_fields", "zone_a_labels"):
        for path in report["bound"]:
            if key in loaded_by_path[path]:
                spec[key] = loaded_by_path[path][key]
                report["sources"][key] = path
                break

    for key, reason in spec_dead_knobs(merged):
        src = report["sources"].get(key, "NOTES_FORMAT.md")
        msg = f"{src}: '{key}' {reason}"
        report["warnings"].append(msg)

    return spec, report


def render_sbc_spec(spec: Dict[str, Any]) -> str:
    """Render the SBC quality-bar prompt text for spec."""
    prompt = spec.get("_prompt", {})
    pieces = [prompt.get("sbc_intro", "")]
    pieces += [f.get("prompt", "") for f in spec.get("sbc_fields", [])]
    body = "\n\n".join(p for p in pieces if p)
    outro = prompt.get("sbc_outro", "")
    return f"{body}\n\n{outro}" if outro else body


def render_zone_a_spec(spec: Dict[str, Any]) -> str:
    """Render the Zone A hook-spec prompt text for spec."""
    prompt = spec.get("_prompt", {})
    intro = prompt.get("zone_a_intro", "")
    archetypes = prompt.get("zone_a_archetypes", [])
    body = intro + "\n" + "\n".join(a for a in archetypes if a) if archetypes else intro
    outro = prompt.get("zone_a_outro", "")
    return f"{body}\n\n{outro}" if outro else body


def spec_labels_list(spec: Dict[str, Any]) -> List[str]:
    """The full label vocabulary treated as SBC field labels."""
    labels = [f"{f['label']}:" for f in spec.get("sbc_fields", [])]
    for extra in (
        "Teacher Guidance:", "Questions and Answers:",
        "Academic and Industry References:", "Recommended Videos:",
        "Key citations:", "Teaching note:", "URL:", "Clip Duration:",
        "Classroom Guidance:",
        "- Classical Foundation:", "- Contemporary Frontier (2024-2026):",
        "- Expected:", "- Guidance:", "- Assessment Link:", "- Key Insight:",
        "Expected:", "Guidance:", "Assessment Link:", "Key Insight:",
        "- Step 1:", "- Step 2:", "- Step 3:", "- Step 4:",
        "Step 1:", "Step 2:", "Step 3:", "Step 4:",
    ):
        if extra not in labels:
            labels.append(extra)
    return labels


def spec_bold_labels(spec: Dict[str, Any]) -> List[str]:
    """Uppercase section labels that get bold-only treatment in Zone A."""
    labels = [lbl.upper() for lbl in spec.get("zone_a_labels", [])]
    for extra in (
        "WARNING:", "CAUTION:", "CONTRAST:", "ANALOGY:", "PUNCHLINE:",
        "FRAMING:", "CASE:", "ADVANTAGES", "DISADVANTAGES",
        "3-STEP ESCALATION:", "VIDEO CUE:", "SCRIPTED EXPLANATION:",
        "🎯 VISUAL OBJECTIVE:", "👁️ GAZE DIRECTION:", "🔄 STEP-BY-STEP WALK-THROUGH:",
        "💡 VISUAL PUNCHLINE:", "VISUAL OBJECTIVE:", "GAZE DIRECTION:",
        "STEP-BY-STEP WALK-THROUGH:", "VISUAL PUNCHLINE:",
    ):
        if extra not in labels:
            labels.append(extra)
    return labels


def spec_font_sizes(spec: Dict[str, Any]) -> Dict[str, int]:
    """Font sizes in hundredths of a point (1000 = 10pt)."""
    fonts = spec.get("fonts", {}) or {}
    return {
        "timing": fonts.get("timing", 1100),
        "header": fonts.get("header", 1000),
        "body": fonts.get("body", 1000),
        "sub": fonts.get("sub", 900),
    }


def split_prompt(prompt: str) -> Tuple[str, str, str]:
    """Split a single-slide prompt into (prefix, slide_block, suffix)."""
    if SLIDE_START not in prompt or SLIDE_END not in prompt:
        raise ValueError(
            "generation prompt is missing structural markers "
            f"({SLIDE_START!r} / {SLIDE_END!r}); refusing to split"
        )
    prefix, rest = prompt.split(SLIDE_START, 1)
    block, suffix = rest.split(SLIDE_END, 1)
    return prefix, block, suffix


SBC_QUALITY_SPEC = render_sbc_spec(DEFAULT_FORMAT_SPEC)
ZONE_A_HOOK_SPEC = render_zone_a_spec(DEFAULT_FORMAT_SPEC)
