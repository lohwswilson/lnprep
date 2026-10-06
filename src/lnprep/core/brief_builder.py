"""
Review Brief Builder for lnprep.

Assembles the review brief handed to the independent reviewer subagent.
House pattern: Assembles a rich markdown brief without making unrequested AI calls.
Includes:
- Resolved NOTES_FORMAT.md specification.
- Deterministic L1 audit result (run in-process).
- Notes text (in full, or flagged-only mode).
- Visual text content (tables, SmartArt, charts) for vision-free fidelity checks.
- Image manifest paths for pixel verification.
- Automated reference check and citation cache status.
- Assessment architecture parsed from the module's LECTURE_NOTES_GUIDE.md.
- Reviewer rubric.
"""

from __future__ import annotations

import os
import re
from pathlib import Path
from typing import Any, Dict, List, Optional, Set, Tuple

from pptx import Presentation

from lnprep.core import citation_db as cc
from lnprep.core import ref_verifier as vr
from lnprep.core.audit_engine import run_audit
from lnprep.core.common import find_guide_file, get_slide_notes_text, load_format_spec, parse_guide
from lnprep.core.pptx_engine import extract_slide_images, extract_slide_visuals


def get_rubric_content() -> str:
    """Load the reviewer rubric markdown."""
    # Only the packaged copy. A machine-local AgentOS path used to be tried second,
    # which put a personal absolute path in the wheel and gave the rubric a silent
    # second source that could drift from the shipped one.
    candidate_paths = [
        Path(__file__).resolve().parent.parent / "resources" / "reviewer-rubric.md",
    ]
    for p in candidate_paths:
        if p.exists():
            try:
                return p.read_text(encoding="utf-8")
            except Exception:
                pass
    return "# Rubric\n\n(Standard Reviewer Rubric: Fidelity, Factuality, Grounding, Pedagogy, Mechanics)"


def parse_range(spec: Optional[str], total: int) -> List[int]:
    """Parse '20-32' or '1,4,9' into a list of 1-indexed slide numbers."""
    if not spec:
        return list(range(1, total + 1))
    out: List[int] = []
    for part in spec.split(","):
        part = part.strip()
        if not part:
            continue
        if "-" in part:
            a, b = part.split("-", 1)
            try:
                out.extend(range(int(a), int(b) + 1))
            except ValueError:
                continue
        else:
            try:
                out.append(int(part))
            except ValueError:
                continue
    return [n for n in out if 1 <= n <= total]


def notes_for(prs: Presentation, wanted: Set[int]) -> Dict[int, str]:
    """Return {slide_num: notes_text} for wanted slide numbers."""
    out: Dict[int, str] = {}
    for i, slide in enumerate(prs.slides):
        num = i + 1
        if num not in wanted:
            continue
        text = get_slide_notes_text(slide)
        if text:
            out[num] = text
    return out


def visual_text(slide: Any, num: int, probe_dir: str) -> Tuple[Optional[List[str]], Optional[str]]:
    """Textual content of a slide's visuals — tables, charts, SmartArt."""
    try:
        v = extract_slide_visuals(slide, num, probe_dir)
    except Exception as exc:
        return None, f"{type(exc).__name__}: {exc}"

    bits: List[str] = []
    for t in v.get("tables", []):
        rows = t.get("data") or []
        flat = []
        for row in rows[:8]:
            cells = row.get("cells") if isinstance(row, dict) else row
            if cells:
                flat.append(" | ".join(str(c) for c in cells))
        if flat:
            bits.append("TABLE:\n      " + "\n      ".join(flat))

    for sa in v.get("smartart", []):
        txt = sa if isinstance(sa, str) else (sa.get("text") if isinstance(sa, dict) else str(sa))
        if txt:
            bits.append(f"SMARTART: {str(txt)[:400]}")

    for c in v.get("charts", []):
        bits.append(f"CHART: {c.get('chart_type', '?')} title={c.get('has_title')}")

    return bits, None


def reference_check_section(A: Any, pptx_path: str, notes_by_slide: Dict[int, str], ttl_days: int) -> None:
    """Embed the automated reference check so reviewer starts from its result."""
    root = cc.find_module_root(pptx_path)
    rep = vr.verify_notes(notes_by_slide, cache_root=root, ttl_days=ttl_days)
    A("## Automated reference check (ref_verifier)")
    A("")
    counts = " · ".join(f"{k}: {v}" for k, v in sorted(rep["counts"].items())) or "no references found"
    A(f"{counts} — **{rep['blocking']} blocking**.")
    A("")
    rows = [f for f in rep["findings"] if f["status"] not in vr.PASSING]
    if rows:
        A("| Slide | Status | Citation / reference | Detail |")
        A("|---|---|---|---|")
        for f in rows:
            who = f.get("citation") or f.get("ref")
            A(f"| {f['slide']} | **{f['status']}** | {str(who)[:70]} | {str(f.get('detail', ''))[:110]} |")
        A("")
    A("Your duties on references: every **blocking** row is a REVISE/FAIL item; for each")
    A("`EXISTS_SUPPORT_UNCONFIRMED` row, open the source and confirm it supports the claim")
    A("(figures, year, attribution); retry `UNREACHABLE` rows by live search. Rows marked")
    A("VERIFIED/CACHED passed a real check — challenge one only if it looks wrong anyway.")
    A("")


def citation_cache_section(A: Any, pptx_path: str, notes_by_slide: Dict[int, str], ttl_days: int) -> None:
    """Append the cache-state audit that drives the reviewer's first job."""
    cits = vr.citations_by_slide(notes_by_slide)
    A("## Citation cache state")
    A("")
    A(f"Module cache TTL: **{ttl_days} days**. This section is the record of what has")
    A("already been verified for this module and what has not.")
    A("")
    if not cits:
        A("- No `(Author, Year)` citations detected in these notes.")
        A("")
        return
    root = cc.find_module_root(pptx_path)
    data = cc.load(root)
    fresh, stale_or_unknown = [], []
    for (author, year), slides in sorted(cits.items()):
        status, entry, left = cc.lookup(data, author, year, ttl_days)
        if status == "HIT_FRESH" and entry is not None:
            fresh.append((author, year, slides, entry, left))
        else:
            stale_or_unknown.append((author, year, slides, status, entry))
    A(f"**{len(cits)} distinct citation(s)** across the reviewed slides: "
      f"{len(fresh)} cached-and-fresh, {len(stale_or_unknown)} requiring live verification.")
    A("")
    if stale_or_unknown:
        A("### Must be live-verified by you (not cached, or TTL expired)")
        A("")
        A("| Citation | Slides | Cache state |")
        A("|---|---|---|")
        for author, year, slides, status, _e in stale_or_unknown:
            sl = ", ".join(str(s) for s in sorted(set(slides)))
            state = "expired (TTL)" if status == "HIT_STALE" else "not cached"
            A(f"| {author} ({year}) | {sl} | **{state}** |")
        A("")
        A("For each: live-search it, confirm it exists and supports the claim, and record")
        A("it with `lnprep cache record` so the next session does not repeat the work.")
        A("")
    if fresh:
        A("### Already verified for this module (served from cache)")
        A("")
        A("| Citation | Slides | Verified | Days left | Evidence |")
        A("|---|---|---|---|---|")
        for author, year, slides, entry, left in fresh:
            sl = ", ".join(str(s) for s in sorted(set(slides)))
            ev = entry.get("doi") or entry.get("url") or entry.get("title") or "—"
            A(f"| {author} ({year}) | {sl} | {(entry.get('verified_at') or '')[:10]} | "
              f"{left:.0f} | {str(ev)[:60]} |")
        A("")
        A("These passed a real verification within the TTL — you do **not** need to")
        A("re-search them. If one looks *wrong anyway*, report it: a stale cache entry is")
        A("a bug worth catching. Do not silently accept a cached citation you doubt.")
        A("")
    A("**Say in your verdict which citations you verified fresh this run and which you")
    A("took from cache.** The audit trail is the point: a citation ships to students,")
    A("so it must always be traceable to a real check with a date.")
    A("")


def build_brief(
    pptx_path: str,
    out_path: Optional[str] = None,
    slides: Optional[str] = None,
    skip_images: bool = False,
    images_dir: Optional[str] = None,
    visual_verify: bool = False,
    flagged_only: bool = False,
    citation_ttl_days: int = 180,
) -> Tuple[str, Dict[str, Any]]:
    """Build the review brief markdown content and summary dictionary."""
    pptx_path = os.path.abspath(pptx_path)
    spec, report = load_format_spec(pptx_path)
    prs = Presentation(pptx_path)
    total = len(prs.slides)
    wanted_slides = parse_range(slides, total)
    wanted_set = set(wanted_slides)

    # Deterministic L1 audit
    audit = run_audit(pptx_path, slide_num=None)

    # Image manifest
    images = None
    if not skip_images:
        try:
            images = extract_slide_images(pptx_path, output_dir=images_dir, verbose=False)
        except Exception as exc:
            images = {"error": f"{type(exc).__name__}: {exc}", "slides": []}

    # Guide context
    guide_path = find_guide_file(pptx_path)
    ctx = parse_guide(guide_path) if guide_path else {}

    # Rubric
    rubric_text = get_rubric_content()

    notes = notes_for(prs, wanted_set)

    lines: List[str] = []
    A = lines.append

    A(f"# Review brief — {os.path.basename(pptx_path)}")
    A("")
    A("You are the **independent reviewer**. You did not write these notes and you")
    A("have not seen the author's reasoning. That is deliberate: a reviewer that")
    A("knows *why* something was written agrees with it.")
    A("")
    A("**You are read-only.** Never write to a deck, never run a write-back script.")
    A("Report findings; the preparer acts on them.")
    A("")
    A(f"- Deck: `{pptx_path}`")
    A(f"- Slides: {total} total; **{len(wanted_set)} in this brief**"
      + (f" ({min(wanted_set)}–{max(wanted_set)})" if wanted_set else ""))
    A(f"- Slides with notes supplied below: {len(notes)}")
    A("- Format spec: "
      + (f"**BOUND** — {', '.join(report['bound'])}" if report.get("bound")
         else ("**DRAFT, NOT BOUND** — " + ", ".join(report["draft"]) if report.get("draft")
               else ("**unreadable** — " + "; ".join(report.get("errors", [])) if report.get("errors")
                     else "none found — built-in default canon"))))
    if report.get("draft"):
        A("  > The module's spec is still a draft, so the notes were prepared against the")
        A("  > **default** canon, not that file. Judge accordingly and say so in your summary.")
    A("")

    # Declared format
    A("## The format the notes are supposed to follow")
    A("")
    fields = [f["label"] for f in spec.get("sbc_fields", [])]
    A("- Zones: " + ", ".join(k for k, v in (spec.get("zones") or {}).items() if v))
    A(f"- SBC header: `{spec.get('sbc_header')}`")
    A(f"- SBC fields (in order): {', '.join('`' + f + ':`' for f in fields)}")
    if spec.get("zone_a_labels"):
        A(f"- Zone A labels: {', '.join('`' + l + ':`' for l in spec['zone_a_labels'])}")
    A(f"- Teacher guidance required: {bool(spec.get('require_teacher_guidance'))}")
    anchors = (spec.get("vocabulary") or {}).get("regional_anchors") or []
    bench = (spec.get("vocabulary") or {}).get("global_benchmarks") or []
    if anchors or bench:
        A(f"- Concrete Example should pair an APAC anchor ({', '.join(anchors[:5])}…) "
          f"with a global benchmark ({', '.join(bench[:4])})")
    A("")

    # L1 audit
    A("## L1 audit result — READ THIS, DO NOT RE-DERIVE IT")
    A("")
    A("`lnprep audit` already decided the mechanical questions below. Spend your")
    A("budget on judgment, not on re-running these.")
    A("")
    A("| Slide | Status | Body | Chrome | Uncovered | Alignment | Quality | Formatting |")
    A("|---|---|---|---|---|---|---|---|")
    for r in audit.get("results", []):
        if r["slide"] not in wanted_set:
            continue
        c = r.get("coverage") or {}
        al = r.get("alignment") or {}
        q = r.get("quality") or {}
        fm = r.get("formatting") or {}
        A(f"| {r['slide']} | {r['status']} | {r.get('body_items','-')} | "
          f"{r.get('chrome_items','-')} | {c.get('missing','-')} | "
          f"{al.get('status','-')} | {q.get('status','-')} | {fm.get('status','-')} |")
    A("")
    A("L1 statuses: `PASS` clean · `GAPS` uncovered body items · `MISALIGNED` SBC "
      "discusses different vocabulary · `WEAK_QUALITY` marker floor failed · "
      "`UNPARSED_SBC` structure could not be parsed, quality bar never ran · "
      "`NO_SBC` no block · `SKIPPED` no substantive body content.")
    A("")
    A("**Slides L1 already flags** — add what the script cannot see, or move on:")
    A("")

    flagged = [r for r in audit.get("results", [])
               if r["slide"] in wanted_set and r["status"] not in ("PASS", "SKIPPED")]
    flagged_nums = {r["slide"] for r in flagged}
    wanted_by_l1 = {r["slide"] for r in audit.get("results", []) if r["slide"] in wanted_set}

    if not flagged:
        A("- (none — every slide in this brief passed L1 structurally)")
    for r in flagged:
        c = r.get("coverage") or {}
        q = r.get("quality") or {}
        A(f"- **Slide {r['slide']}** `{r['status']}` — {r.get('title','')[:60]}")
        for it in (c.get("missing_items") or [])[:3]:
            A(f"    - uncovered: {it[:90]}")
        if q.get("reason"):
            A(f"    - {q['reason']}")
    A("")

    # Automated reference check
    try:
        reference_check_section(A, pptx_path, notes, citation_ttl_days)
    except Exception as exc:
        A("## Automated reference check (ref_verifier)")
        A("")
        A(f"- Check unavailable ({exc.__class__.__name__}: {exc}). Verify **every** reference by live")
        A("  search, including Concrete Example sources and their figures.")
        A("")

    # Citation cache state
    try:
        citation_cache_section(A, pptx_path, notes, citation_ttl_days)
    except Exception as exc:
        A("## Citation cache state")
        A("")
        A(f"- Cache unavailable ({exc.__class__.__name__}: {exc}). Verify **every** citation by live")
        A("  search as normal; nothing may be served from cache this run.")
        A("")

    # Images
    A("## Visuals — open these to check visual fidelity (rubric check 1)")
    A("")
    if images is None:
        A("- (image extraction skipped)")
    elif images.get("error"):
        A(f"- (image extraction failed: {images['error']})")
    elif not images.get("slides_with_visuals"):
        A("- (no extractable images or charts in this deck)")
    else:
        A(f"- {images.get('images_extracted', 0)} image(s) across "
          f"{images.get('slides_with_visuals', 0)} slide(s) → `{images.get('output_dir', '')}/`")
        A("")
        for s in images.get("slides", []):
            if s["slide"] not in wanted_set or not s.get("visuals"):
                continue
            bits = []
            for v in s["visuals"]:
                if v.get("type") == "image" and v.get("path"):
                    bits.append(f"`{v['path']}` ({v.get('size','?')})")
                elif v.get("type") == "chart":
                    bits.append(f"chart {v.get('chart_type','?')}")
                else:
                    bits.append(v.get("type", "unknown"))
            A(f"- Slide {s['slide']}: " + "; ".join(bits))
    A("")

    # Visual text
    probe_dir = os.path.join((images or {}).get("output_dir") or "/tmp", "_text_probe")
    A("## What the visuals SAY (text — works without image input)")
    A("")
    A("Compare the notes against this. If the notes describe a diagram whose table")
    A("rows or SmartArt labels are nothing like these, that is a visual mismatch.")
    A("This cannot judge layout or aesthetics — only content.")
    A("")
    any_vt = False
    for i, slide in enumerate(prs.slides):
        num = i + 1
        if num not in wanted_set:
            continue
        vbits, err = visual_text(slide, num, probe_dir)
        if err:
            A(f"- Slide {num}: (visual text unavailable: {err})")
            continue
        if vbits:
            any_vt = True
            A(f"- **Slide {num}**:")
            for b in vbits:
                A(f"    - {b}")
    if not any_vt:
        A("- (no tables, SmartArt or charts in this range)")
    A("")

    # Visual fidelity verdicts
    A("## Visual fidelity verdicts")
    A("")
    if not visual_verify:
        A("- **Not run.** Pass `--visual-verify` to run vision-model comparisons.")
        A("  Without it, judge visual fidelity only from the text above, and state")
        A("  in your summary that pixel-level fidelity was not checked.")
        A("")
    else:
        A("- Automated vision verification selected.")
        A("")

    # Assessment architecture
    if ctx:
        A("## Assessment architecture (from the module's guide)")
        A("")
        A("Use this to check `Assessment Link:` claims against the real tasks and word")
        A("counts, rather than accepting them at face value.")
        A("")
        for k in ("module_code", "institution", "term", "session", "duration", "style"):
            if ctx.get(k):
                A(f"- {k}: {ctx[k]}")
        for k, v in ctx.items():
            if "assessment" in k.lower() or "mlo" in k.lower():
                A(f"- {k}: {str(v)[:200]}")
        A("")

    # Notes under review
    A("## The notes under review")
    A("")
    if flagged_only:
        full_nums = sorted(n for n in notes if n in flagged_nums or n not in wanted_by_l1)
        stub_nums = sorted(n for n in notes if n not in full_nums)
        if not stub_nums:
            A("**Compressed mode requested (`--flagged-only`), but it had no effect:** "
              "every slide in this brief is L1-flagged, so all note text is supplied "
              "in full.")
            A("")
        else:
            A(f"**Compressed mode (`--flagged-only`).** Full note text is supplied for "
              f"**{len(full_nums)}** slide(s). The other **{len(stub_nums)}** slide(s) "
              f"passed L1 structurally, so only a summary is given.")
            A("")
        for num in full_nums:
            A(f"### Slide {num}")
            A("")
            A("```")
            A(notes[num].strip())
            A("```")
            A("")
        if stub_nums:
            A("### Slides passing L1 (summary only — full text omitted)")
            A("")
            for num in stub_nums:
                body = notes[num].strip()
                title = ""
                for ln in body.splitlines():
                    if ln.strip():
                        title = ln.strip()[:110]
                        break
                anchors_found = re.findall(r"\*\s*MLO Anchor:\s*(.+)", body)
                first_line = ""
                for ln in body.splitlines():
                    s = ln.strip()
                    if s and not s.startswith(("---", "⏱", body.splitlines()[0].strip() if body else "")):
                        first_line = s[:150]
                        break
                A(f"- **Slide {num}** — {title}")
                if anchors_found:
                    A(f"    - MLO anchor: {anchors_found[0][:90]}")
                if first_line:
                    A(f"    - opens: {first_line}")
            A("")
    else:
        for num in sorted(notes):
            A(f"### Slide {num}")
            A("")
            A("```")
            A(notes[num].strip())
            A("```")
            A("")

    A("---")
    A("")
    A("# Rubric")
    A("")
    A(rubric_text)

    text = "\n".join(lines)
    stats = {
        "chars": len(text),
        "slides": len(notes),
        "l1_flagged": len(flagged),
        "images": (images or {}).get("images_extracted", 0) if images else 0,
    }
    if out_path:
        out_p = Path(out_path).resolve()
        out_p.parent.mkdir(parents=True, exist_ok=True)
        out_p.write_text(text, encoding="utf-8")
        stats["out"] = str(out_p)

    return text, stats
