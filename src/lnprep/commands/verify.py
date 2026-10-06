"""
lnprep verify command: Verify academic citations, DOIs, ISBNs, and URLs in slide notes.
"""

from __future__ import annotations

import json
from pathlib import Path
from typing import Dict, Optional, Set

import typer
from pptx import Presentation
from rich.table import Table

from lnprep.console import console, print_error, print_formatted, print_info, print_success
from lnprep.core import citation_db as cc
from lnprep.core import ref_verifier as vr
from lnprep.core.common import get_slide_notes_text

app = typer.Typer(help="Verify citations, evidence links, and Concrete Example sources")


def _parse_range(spec: Optional[str], total: int) -> Set[int]:
    if not spec or spec.strip().lower() == "all":
        return set(range(1, total + 1))
    out: Set[int] = set()
    for part in spec.split(","):
        part = part.strip()
        if not part:
            continue
        if "-" in part:
            a, b = part.split("-", 1)
            try:
                out.update(range(int(a), int(b) + 1))
            except ValueError:
                continue
        else:
            try:
                out.add(int(part))
            except ValueError:
                continue
    return {n for n in out if 1 <= n <= total}


@app.callback(invoke_without_command=True)
def verify_command(
    ctx: typer.Context,
    pptx_path: Path = typer.Argument(
        ...,
        help="Path to PPTX file",
        exists=True,
        file_okay=True,
        dir_okay=False,
        resolve_path=True,
    ),
    slides: Optional[str] = typer.Option(
        None,
        "--slides",
        "-s",
        help="Slide range, e.g. '1-10' or '3,5,8'",
    ),
    notes_json: Optional[Path] = typer.Option(
        None,
        "--notes-json",
        help="Path to JSON file with notes {slide: text}",
    ),
    text_file: Optional[Path] = typer.Option(
        None,
        "--text",
        help="Path to single text file containing notes (tested as Slide 1)",
    ),
    offline: bool = typer.Option(
        False,
        "--offline",
        help="Offline mode — check structure without network calls",
    ),
    no_record: bool = typer.Option(
        False,
        "--no-record",
        help="Do not record verified citations to the module cache",
    ),
    ttl_days: int = typer.Option(
        180,
        "--ttl-days",
        help="TTL for citation cache in days",
    ),
    as_json: bool = typer.Option(
        False,
        "--json",
        help="Output findings as JSON",
    ),
) -> None:
    """Verify citations, DOIs, ISBNs, and URLs in slide notes."""
    notes_by_slide: Dict[int, str] = {}

    if text_file:
        content = Path(text_file).read_text(encoding="utf-8")
        notes_by_slide[1] = content
    elif notes_json:
        with open(notes_json, "r", encoding="utf-8") as f:
            data = json.loads(f.read())
        if isinstance(data, list):
            for item in data:
                notes_by_slide[int(item["slide"])] = item.get("notes", "")
        elif isinstance(data, dict):
            for k, v in data.items():
                notes_by_slide[int(k)] = v if isinstance(v, str) else v.get("notes", "")
    else:
        prs = Presentation(str(pptx_path))
        total = len(prs.slides)
        wanted = _parse_range(slides, total)
        for i, slide in enumerate(prs.slides):
            num = i + 1
            if num in wanted:
                txt = get_slide_notes_text(slide)
                if txt:
                    notes_by_slide[num] = txt

    root = cc.find_module_root(str(pptx_path))
    rep = vr.verify_notes(
        notes_by_slide,
        offline=offline,
        cache_root=root,
        record_new=not no_record,
        ttl_days=ttl_days,
    )

    if as_json:
        print_formatted(rep, as_json=True)
        if rep.get("blocking", 0) > 0:
            raise typer.Exit(code=1)
        return

    mode_str = " (OFFLINE — structural only)" if rep.get("offline") else ""
    console.print(f"[bold cyan]Reference Verification{mode_str}:[/bold cyan] {pptx_path.name}")

    counts = rep.get("counts", {})
    count_summary = " · ".join(f"{k}: {v}" for k, v in sorted(counts.items())) or "no references found"
    console.print(f"Summary: {count_summary}\n")

    findings = rep.get("findings", [])
    if findings:
        table = Table(title="Reference Findings", show_header=True)
        table.add_column("Slide", style="cyan", width=6)
        table.add_column("Status", style="bold", width=22)
        table.add_column("Citation / Target", max_width=35)
        table.add_column("Detail", max_width=45)

        status_colors = {
            "VERIFIED": "green",
            "CACHED": "green",
            "EXISTS_SUPPORT_UNCONFIRMED": "yellow",
            "UNREACHABLE": "dim",
            "MISMATCH": "red",
            "NOT_FOUND": "red",
            "BROKEN_LINK": "red",
            "EPHEMERAL_URL": "red",
            "NO_SOURCE": "red",
            "UNRESOLVED": "red",
        }

        for f in findings:
            status = f.get("status", "UNKNOWN")
            color = status_colors.get(status, "white")
            who = f.get("citation") or f.get("ref") or "—"
            table.add_row(
                str(f.get("slide", "-")),
                f"[{color}]{status}[/{color}]",
                str(who)[:35],
                str(f.get("detail", ""))[:45],
            )

        console.print(table)
    else:
        print_info("No citations or references found to verify.")

    if rep.get("recorded_to_cache"):
        console.print(f"\n[green]Recorded {rep['recorded_to_cache']} verified citation(s) to module cache.[/green]")

    blocking = rep.get("blocking", 0)
    if blocking > 0:
        print_error(f"BLOCKING: {blocking} finding(s) must be resolved before write-back!")
        raise typer.Exit(code=1)
    else:
        print_success("Verification passed: 0 blocking issues.")
