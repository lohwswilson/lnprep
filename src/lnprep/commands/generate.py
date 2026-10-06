"""
lnprep generate command: Generate structured lecture note prompts for PPTX slides.
"""

from __future__ import annotations

import json
import re
from pathlib import Path
from typing import List, Optional

import typer
from pptx import Presentation

from lnprep.console import console, print_error, print_formatted, print_success
from lnprep.core import citation_db as cc
from lnprep.core.common import find_guide_file, load_format_spec, parse_guide
from lnprep.core.pptx_engine import extract_slide_context
from lnprep.core.prompt_engine import build_batch_prompt, build_generation_prompt, parse_slide_spec

app = typer.Typer(help="Generate lecture note prompts for slides")


@app.callback(invoke_without_command=True)
def generate_command(
    ctx: typer.Context,
    pptx_path: Path = typer.Argument(
        ...,
        help="Path to PPTX file",
        exists=True,
        file_okay=True,
        dir_okay=False,
        resolve_path=True,
    ),
    slide: Optional[int] = typer.Option(
        None,
        "--slide",
        "-s",
        help="Generate prompt for a single slide (1-indexed)",
    ),
    batch: Optional[str] = typer.Option(
        None,
        "--batch",
        "-b",
        help="Generate ONE batched prompt for slide spec, e.g. '1,5-12' or 'all'",
    ),
    from_gap: Optional[Path] = typer.Option(
        None,
        "--from-gap",
        help="Generate prompt for slides flagged in gap scan JSON",
    ),
    session: Optional[int] = typer.Option(
        None,
        "--session",
        help="Session number for context",
    ),
    prompt_only: bool = typer.Option(
        True,
        "--prompt-only",
        help="Output prompt only (lnprep follows house pattern: no unrequested AI calls)",
    ),
    out: Optional[Path] = typer.Option(
        None,
        "--out",
        "-o",
        help="Write prompt to output file",
    ),
    as_json: bool = typer.Option(
        False,
        "--json",
        help="Output context and prompt as JSON",
    ),
) -> None:
    """Generate structured lecture note prompt for slides."""
    prs = Presentation(str(pptx_path))
    total_slides = len(prs.slides)

    guide_path = find_guide_file(str(pptx_path))
    module_context = parse_guide(guide_path) if guide_path else {}
    format_spec, format_report = load_format_spec(str(pptx_path))

    module_root = cc.find_module_root(str(pptx_path))
    citation_data = cc.load(module_root)
    cached_citations = citation_data if citation_data and citation_data.get("citations") else None

    session_info = None
    if session and module_context.get("_session_topics_raw"):
        sessions_block = module_context["_session_topics_raw"]
        session_pattern = re.compile(
            rf'-\s+number:\s+{session}\s*\n\s+date:\s*"([^"]+)"\s*\n\s+format:\s*"([^"]+)"\s*\n\s+topics:\s*\n((?:\s+-.*\n?)*)',
            re.MULTILINE,
        )
        match = session_pattern.search(sessions_block)
        if match:
            topics = re.findall(r'-\s*"([^"]+)"', match.group(3))
            session_info = {
                "number": session,
                "date": match.group(1),
                "format": match.group(2),
                "topics": topics,
            }

    slides_to_process: List[int] = []
    if batch:
        try:
            parsed = parse_slide_spec(batch, total_slides)
            slides_to_process = parsed if parsed else []
        except ValueError as e:
            print_error(f"Invalid --batch spec '{batch}': {e}")
            raise typer.Exit(code=1)
        if not slides_to_process:
            print_error(f"--batch '{batch}' selected no slides")
            raise typer.Exit(code=1)
    elif slide:
        if not (1 <= slide <= total_slides):
            print_error(f"Slide {slide} out of range (1..{total_slides})")
            raise typer.Exit(code=1)
        slides_to_process = [slide]
    elif from_gap:
        try:
            with open(from_gap, "r", encoding="utf-8") as f:
                gap_data = json.loads(f.read(), strict=False)
            slides_to_process = [
                s["slide"]
                for s in gap_data.get("slides", [])
                if s.get("status") in ("empty", "minimal") and 1 <= s.get("slide", 0) <= total_slides
            ]
        except Exception as e:
            print_error(f"Failed to read --from-gap file: {e}")
            raise typer.Exit(code=1)
        if not slides_to_process:
            print_error("No empty or minimal slides found in gap file")
            raise typer.Exit(code=1)
    else:
        print_error("Specify --slide N, --batch <spec>, or --from-gap <file>")
        raise typer.Exit(code=1)

    slide_contexts = [
        extract_slide_context(prs.slides[s - 1], s)
        for s in slides_to_process
    ]

    if len(slide_contexts) == 1:
        prompt_text = build_generation_prompt(
            slide_contexts[0],
            module_context,
            session_info=session_info,
            spec=format_spec,
            cached_citations=cached_citations,
        )
    else:
        prompt_text = build_batch_prompt(
            slide_contexts,
            module_context,
            session_info=session_info,
            spec=format_spec,
            cached_citations=cached_citations,
        )

    if out:
        out_p = Path(out).resolve()
        out_p.parent.mkdir(parents=True, exist_ok=True)
        out_p.write_text(prompt_text, encoding="utf-8")
        if not as_json:
            print_success(f"Prompt saved to {out_p} ({len(prompt_text)} chars, {len(slides_to_process)} slides)")

    if as_json:
        print_formatted(
            {
                "slides": slides_to_process,
                "prompt": prompt_text,
                "out": str(out) if out else None,
                "chars": len(prompt_text),
            },
            as_json=True,
        )
    elif not out:
        console.print(prompt_text)
