"""
lnprep enhance command: Generate targeted prompts to enhance partial or deficient slide notes.
"""

from __future__ import annotations

import json
from pathlib import Path
from typing import List, Optional

import typer
from pptx import Presentation

from lnprep.console import console, print_error, print_formatted, print_success
from lnprep.core import citation_db as cc
from lnprep.core.common import find_guide_file, get_slide_notes_text, load_format_spec, parse_guide
from lnprep.core.prompt_engine import audit_notes_structure, build_enhancement_prompt

app = typer.Typer(help="Generate targeted enhancement prompts for incomplete slide notes")


@app.callback(invoke_without_command=True)
def enhance_command(
    ctx: typer.Context,
    pptx_path: Path = typer.Argument(
        ...,
        help="Path to PPTX file",
        exists=True,
        file_okay=True,
        dir_okay=False,
        resolve_path=True,
    ),
    slide: Optional[int] = typer.Option(
        None,
        "--slide",
        "-s",
        help="Target slide number (1-indexed)",
    ),
    from_gap: Optional[Path] = typer.Option(
        None,
        "--from-gap",
        help="Path to gap scan JSON file",
    ),
    prompt_only: bool = typer.Option(
        True,
        "--prompt-only",
        help="Output prompt only",
    ),
    out: Optional[Path] = typer.Option(
        None,
        "--out",
        "-o",
        help="Write enhancement prompt to file",
    ),
    as_json: bool = typer.Option(
        False,
        "--json",
        help="Output JSON summary and prompt",
    ),
) -> None:
    """Generate targeted enhancement prompt for slides with missing structural elements."""
    prs = Presentation(str(pptx_path))
    total_slides = len(prs.slides)

    guide_path = find_guide_file(str(pptx_path))
    module_context = parse_guide(guide_path) if guide_path else {}
    format_spec, _ = load_format_spec(str(pptx_path))

    module_root = cc.find_module_root(str(pptx_path))
    citation_data = cc.load(module_root)
    cached_citations = citation_data if citation_data and citation_data.get("citations") else None

    slides_to_process: List[int] = []
    if slide:
        if not (1 <= slide <= total_slides):
            print_error(f"Slide {slide} out of range (1..{total_slides})")
            raise typer.Exit(code=1)
        slides_to_process = [slide]
    elif from_gap:
        try:
            with open(from_gap, "r", encoding="utf-8") as f:
                gap_data = json.loads(f.read(), strict=False)
            slides_to_process = [
                s["slide"]
                for s in gap_data.get("slides", [])
                if s.get("status") == "partial" and 1 <= s.get("slide", 0) <= total_slides
            ]
        except Exception as e:
            print_error(f"Failed to read gap file: {e}")
            raise typer.Exit(code=1)
        if not slides_to_process:
            print_error("No partial slides found in gap file")
            raise typer.Exit(code=1)
    else:
        print_error("Specify --slide N or --from-gap <file>")
        raise typer.Exit(code=1)

    all_prompts = []
    results = []

    for s_num in slides_to_process:
        s_obj = prs.slides[s_num - 1]
        existing_notes = get_slide_notes_text(s_obj)
        audit_info = audit_notes_structure(existing_notes)

        p = build_enhancement_prompt(
            s_obj,
            s_num,
            existing_notes,
            audit_info,
            module_context,
            spec=format_spec,
            cached_citations=cached_citations,
        )
        all_prompts.append(f"### Slide {s_num}\n\n{p}")
        results.append({
            "slide": s_num,
            "missing_count": len(audit_info.get("missing", [])),
            "missing": audit_info.get("missing", []),
            "prompt": p,
        })

    combined_text = "\n\n---\n\n".join(all_prompts)

    if out:
        out_p = Path(out).resolve()
        out_p.parent.mkdir(parents=True, exist_ok=True)
        out_p.write_text(combined_text, encoding="utf-8")
        if not as_json:
            print_success(f"Enhancement prompt saved to {out_p}")

    if as_json:
        print_formatted({"results": results, "out": str(out) if out else None}, as_json=True)
    elif not out:
        console.print(combined_text)
