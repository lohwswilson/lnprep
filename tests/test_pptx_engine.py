"""Tests for PPTX extraction — group walking and alt text.

This module had no direct test file; both behaviours below were silent content
losses rather than crashes, so nothing surfaced them.
"""

from PIL import Image
from pptx import Presentation
from pptx.util import Inches

from lnprep.core.pptx_engine import extract_slide_context, extract_slide_images, shape_alt_text


def _tag_alt(picture, text):
    for node in picture._element.iter():
        if node.tag.endswith("}cNvPr"):
            node.set("descr", text)
            return


def _deck_with_nested_picture(tmp_path):
    Image.new("RGB", (60, 40), "red").save(str(tmp_path / "red.png"))
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "Grouped content"

    top = slide.shapes.add_picture(str(tmp_path / "red.png"), Inches(1), Inches(1))
    _tag_alt(top, "ALT TEXT: top level")

    nested = slide.shapes.add_picture(str(tmp_path / "red.png"), Inches(3), Inches(1))
    _tag_alt(nested, "ALT TEXT: nested in group")
    group = slide.shapes.add_group_shape()
    group.shapes._spTree.append(nested._element)

    path = str(tmp_path / "deck.pptx")
    prs.save(path)
    return path


def test_extract_slide_context_walks_into_groups(tmp_path):
    """Grouped shapes used to be invisible to prompt generation.

    The context extractor iterated slide.shapes only and represented a group as
    nothing but its name, so a nested picture or table never reached the prompt —
    even though the audit engine's own walker saw it.
    """
    path = _deck_with_nested_picture(tmp_path)
    prs = Presentation(path)
    context = extract_slide_context(prs.slides[0], 1)

    alts = [v.get("alt_text") for v in context["visuals"] if v.get("alt_text")]
    assert len(alts) == 2, f"expected both pictures, got {alts}"
    assert "nested in group" in alts[1]


def test_alt_text_reaches_the_image_manifest(tmp_path):
    """Alt text was never read, though it is often the only description of a diagram."""
    path = _deck_with_nested_picture(tmp_path)
    manifest = extract_slide_images(path, output_dir=str(tmp_path / "imgs"))

    alts = [
        v.get("alt_text")
        for slide in manifest["slides"]
        for v in slide.get("visuals", [])
        if v.get("alt_text")
    ]
    assert len(alts) == 2


def test_shape_alt_text_is_safe_on_shapes_without_one():
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(2), Inches(1))
    assert shape_alt_text(box) == ""
