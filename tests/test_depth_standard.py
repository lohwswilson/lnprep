"""Tests for SBC and Zone A depth standards."""

import os

from pptx import Presentation
from pptx.util import Inches

from lnprep.core import audit_engine as sa

FIELD_LABELS = [
    "Plain English",
    "Deep Research",
    "Concrete Example",
    "Bigger Picture",
    "Assessment Link",
    "Manager's So What",
]
CFG = {"field_labels": FIELD_LABELS, "min_paragraphs": 2}

SHALLOW = """Stakeholders:

Plain English: A stakeholder is anyone affected by the operation.

Deep Research:
- Classical Foundation: Freeman (1984) is a study of stakeholders.
- Contemporary Frontier (2024-2026): Recent research describes stakeholders.

Concrete Example: For example, PSA Tuas Port serves stakeholders.

Bigger Picture: Strategically, stakeholders matter for consequences.

Assessment Link: This applies in the assignment report.

Manager's So What: Actionable takeaway: managers should act.
"""

_KW = {
    "Plain English": "A stakeholder is anyone affected by the operation.",
    "Deep Research": "Freeman (1984) is the seminal study of this mechanism.",
    "Concrete Example": "For example, PSA Tuas Port serves its stakeholders.",
    "Bigger Picture": "Strategically, the consequence for the network matters.",
    "Assessment Link": "This applies in the assignment report you will write.",
    "Manager's So What": "Actionable takeaway: managers should act on this.",
}
_FILL = (
    "This second sentence exists only to clear the fragment floor so the "
    "paragraph is counted rather than discarded as stray text."
)
DEEP = (
    "Stakeholders:\n\n"
    + "\n\n".join(f"{lab}: {_KW[lab]} {_FILL}\n\n{_KW[lab]} {_FILL}" for lab in FIELD_LABELS)
    + "\n"
)


def test_keyword_vs_depth():
    q_shallow = sa.audit_sbc_quality(SHALLOW, CFG)
    q_deep = sa.audit_sbc_quality(DEEP, CFG)
    assert q_shallow["average_score"] == q_deep["average_score"]

    d_shallow = sa.audit_sbc_depth(SHALLOW, CFG)
    d_deep = sa.audit_sbc_depth(DEEP, CFG)
    assert d_shallow["status"] == "BELOW_DEPTH"
    assert d_deep["status"] == "DEPTH_OK"
    assert len(d_shallow["below"]) == 5
    assert "Plain English" not in d_shallow["below"]
    assert d_deep["average_paragraphs"] >= 2.0
    assert d_deep["fields_checked"] == 6


def test_paragraph_counting_blank_lines():
    one_blob = "Stakeholders:\n\nBigger Picture: " + ("This sentence has plenty of words. " * 20)
    r = sa.audit_sbc_depth(one_blob, CFG)
    assert r["items"][0]["paragraphs"] == 1
    assert r["status"] == "BELOW_DEPTH"
    assert r["items"][0]["chars"] > 500


def test_zone_a_depth():
    kp_one = '* KEY POINT: "A single paragraph explanation of the core principle here."'
    r1 = sa.audit_zone_a_depth(kp_one, CFG)
    assert r1["status"] == "DEPTH_OK"

    kp_two = (
        '* KEY POINT: "First paragraph explaining the core operational principle '
        'in enough detail to clear the fragment floor and mean something."\n\n'
        '"Second paragraph covering the operational mechanism, system friction and '
        'strategic consequence for the assessment."'
    )
    r2 = sa.audit_zone_a_depth(kp_two, CFG)
    assert r2["status"] == "DEPTH_OK"
    assert r2["average_paragraphs"] >= 2.0


def test_advisory_boundary():
    d_shallow = sa.audit_sbc_depth(SHALLOW, CFG)
    assert d_shallow["status"] in ("BELOW_DEPTH", "DEPTH_OK", "NO_FIELDS_PARSED", "NO_SBC")

    d3 = sa.audit_sbc_depth(DEEP, {"field_labels": FIELD_LABELS, "min_paragraphs": 3})
    assert d3["status"] == "BELOW_DEPTH"

    d_def = sa.audit_sbc_depth(DEEP, {"field_labels": FIELD_LABELS})
    assert d_def["min_paragraphs"] == 2


def test_degenerate_input():
    assert sa.audit_sbc_depth("", CFG)["status"] == "NO_SBC"
    assert sa.audit_sbc_depth("just prose with no labels at all here", CFG)["status"] in (
        "NO_FIELDS_PARSED",
        "BELOW_DEPTH",
    )
    assert sa.audit_zone_a_depth("", CFG)["status"] == "NO_ZONE_A"
    assert sa.audit_zone_a_depth('BRIDGE: "hello"', CFG)["status"] == "NO_KEY_POINT"


def test_prose_rules():
    _P = (
        "This paragraph is long enough to clear the forty character fragment floor "
        "used by the paragraph counter."
    )
    kp_four = '* KEY POINT: "' + "\n\n".join([_P] * 4) + '"'
    r4 = sa.audit_zone_a_depth(kp_four, CFG)
    assert r4["status"] == "OVER_MAX_PARAGRAPHS"

    kp_three = '* KEY POINT: "' + "\n\n".join([_P] * 3) + '"'
    assert sa.audit_zone_a_depth(kp_three, CFG)["status"] == "DEPTH_OK"

    kp_bullets = '* KEY POINT: "' + _P + '\n- first bullet point here\n- second bullet point here"'
    rb = sa.audit_zone_a_depth(kp_bullets, CFG)
    assert rb["status"] == "NOT_PROSE"
    assert rb["min_paragraphs"] == 1 and rb["max_paragraphs"] == 3

    pe_four = "Item:\n\nPlain English: " + "\n\n".join([_P] * 4)
    d = sa.audit_sbc_depth(pe_four, CFG)
    assert d["over"] == ["Plain English"]

    pe_bullets = "Item:\n\nPlain English: " + _P + "\n• a bullet that should be prose\n• another one"
    d_b = sa.audit_sbc_depth(pe_bullets, CFG)
    assert d_b["not_prose"] == ["Plain English"]

    ce = (
        "Item:\n\nConcrete Example: For example, "
        + _P
        + '\n\nSources: Org (2025) "Title of the source document here", https://example.org/a-long-enough-url'
    )
    d_ce = sa.audit_sbc_depth(ce, CFG)
    assert d_ce["items"][0]["paragraphs"] == 1


def _deck_with_notes(directory: str, notes: str) -> str:
    path = os.path.join(directory, "deck.pptx")
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "Risk: Bullwhip Effect"
    box = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(6), Inches(2))
    box.text_frame.text = "Bullwhip Effect amplification across tiers"
    slide.notes_slide.notes_text_frame.text = notes
    prs.save(path)
    return path


def test_verdict_fails_a_slide_whose_sbc_fields_are_stubs(tmp_path):
    """Depth and example sourcing must be able to fail the verdict.

    Regression: audit_slide computed depth, Zone A depth and example sourcing and
    then ignored all three, so a slide whose six SBC fields each held a single
    word was reported as PASS.
    """
    notes = (
        "--- SPEAKER NOTES ---\nTiming: 5 min\n\n"
        "KEY POINT: A short but valid key point sentence here.\n\n"
        "--- VISUAL DECONSTRUCTION ---\n\n"
        "--- LECTURE NOTES ---\n"
        "• SLIDE BODY COVERAGE — RISK:\n"
        "Bullwhip Effect amplification across tiers:\n"
        "Plain English: is\nDeep Research: research\nConcrete Example: example\n"
        "Bigger Picture: x\nAssessment Link: y\nManager's So What: z"
    )
    result = sa.run_audit(_deck_with_notes(str(tmp_path), notes))["results"][0]

    assert result["depth"]["status"] == "BELOW_DEPTH"
    assert result["example_sources"]["status"] == "EXAMPLE_UNSOURCED"
    assert result["status"] == "SHALLOW_DEPTH"


def test_run_audit_reports_a_status_summary(tmp_path):
    """run_audit must return a 'summary' its callers can render.

    Regression: the key did not exist, so `lnprep audit` printed a table of zeros
    for every status and `lnprep check`'s audit half was unconditionally zero --
    the composite gate passed decks with no notes at all.
    """
    deck = _deck_with_notes(str(tmp_path), "--- LECTURE NOTES ---\nCore Narrative: x\n")
    res = sa.run_audit(deck)

    assert "summary" in res
    assert sum(res["summary"].values()) == len(res["results"]) == 1
    assert res["summary"]["no_sbc"] == 1
    assert res["summary"]["pass"] == 0


def test_summary_counts_every_documented_status():
    """The summary must carry a bucket for every status audit_slide can emit."""
    summary = sa.summarise_results(
        [{"status": s} for s in sa.AUDIT_STATUSES] + [{"status": "PASS"}]
    )
    assert set(summary) == {s.lower() for s in sa.AUDIT_STATUSES}
    assert summary["pass"] == 2
    assert sum(summary.values()) == len(sa.AUDIT_STATUSES) + 1
