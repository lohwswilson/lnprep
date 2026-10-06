"""Tests for brief_builder module."""

import os
import shutil
import tempfile

import pytest
from pptx import Presentation

from lnprep.core import brief_builder as bb
from lnprep.core.writer_engine import update_slide


@pytest.fixture
def sample_presentation():
    tmp_dir = tempfile.mkdtemp(prefix="lnprep-test-brief-")
    deck_path = os.path.join(tmp_dir, "brief_test_deck.pptx")
    prs = Presentation()
    slide_layout = prs.slide_layouts[1]
    s1 = prs.slides.add_slide(slide_layout)
    s1.shapes.title.text = "Strategic Trade-offs"
    prs.save(deck_path)

    # Write notes to slide 1
    sample_notes = """--- SPEAKER NOTES ---
⏱ 4.0 min | 🎯 Trade-off Analysis
BRIDGE: "Welcome back."
KEY POINT: "Trade-offs are unavoidable."

--- LECTURE NOTES ---
• CORE NARRATIVE:
* MLO Anchor: [MLO 1]
"""
    update_slide(deck_path, 1, sample_notes)

    yield deck_path
    shutil.rmtree(tmp_dir, ignore_errors=True)


def test_parse_range():
    assert bb.parse_range("1-5", 10) == [1, 2, 3, 4, 5]
    assert bb.parse_range("2,4,6", 10) == [2, 4, 6]
    assert bb.parse_range(None, 3) == [1, 2, 3]


def test_build_brief(sample_presentation):
    text, stats = bb.build_brief(sample_presentation, skip_images=True)
    assert stats["slides"] >= 1
    assert "Review brief —" in text
    assert "## L1 audit result" in text
    assert "## The notes under review" in text
    assert "# Rubric" in text
