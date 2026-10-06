"""Tests for writer_engine direct XML injection, backup, and payload parsing."""

import os
import shutil
import tempfile
import pytest
from pptx import Presentation

from lnprep.core import writer_engine as we
from lnprep.core.common import get_slide_notes_text


def _notes(text: str = "A valid key point sentence for testing.") -> str:
    return (
        "--- SPEAKER NOTES ---\n"
        f"KEY POINT: {text}\n\n"
        "--- LECTURE NOTES ---\n"
        "Core Narrative: Body text.\n"
    )


def _write_deck(path: str, titles) -> None:
    prs = Presentation()
    for title in titles:
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        slide.shapes.title.text = title
    prs.save(path)


@pytest.fixture
def sample_deck():
    tmp_dir = tempfile.mkdtemp(prefix="lnprep-test-deck-")
    deck_path = os.path.join(tmp_dir, "test_deck.pptx")
    prs = Presentation()
    slide_layout = prs.slide_layouts[1]
    slide1 = prs.slides.add_slide(slide_layout)
    slide1.shapes.title.text = "Introduction"
    slide2 = prs.slides.add_slide(slide_layout)
    slide2.shapes.title.text = "Core Principles"
    prs.save(deck_path)

    yield deck_path
    shutil.rmtree(tmp_dir, ignore_errors=True)


def test_load_notes_payload():
    valid = '{"slides": [{"slide": 1, "notes": "Test notes"}]}'
    res = we.load_notes_payload(valid)
    assert res["slides"][0]["slide"] == 1

    fenced = "```json\n" + valid + "\n```"
    res_f = we.load_notes_payload(fenced)
    assert res_f["slides"][0]["slide"] == 1

    with pytest.raises(ValueError):
        we.load_notes_payload("this is not json at all")


def test_backup_once(sample_deck):
    bak = we.backup_once(sample_deck)
    assert os.path.exists(bak)
    assert "_backup_" in bak
    # Calling again returns same backup without overwriting
    bak2 = we.backup_once(sample_deck)
    assert bak == bak2


def test_update_slide_xml_injection(sample_deck):
    notes = """--- SPEAKER NOTES ---
⏱ 3.0 min | 🎯 Supply Chain Friction

BRIDGE: "Linking from inventory principles"
NARRATIVE OPENING / HOOK: "When the Suez Canal was blocked, shipping rates tripled overnight."

KEY POINT: "Operational bottlenecks propagate non-linearly across interconnected networks."

--- VISUAL DECONSTRUCTION ---
🎯 VISUAL OBJECTIVE:
Map the ripple effect across European port terminals.

👁️ GAZE DIRECTION:
Look first at the horizontal axis.

🔄 STEP-BY-STEP WALK-THROUGH:
- Step 1: Trace the initial blockage at point A.
- Step 2: Observe queue build-up at port B.

💡 VISUAL PUNCHLINE:
"Single points of failure dictate network capacity."

--- LECTURE NOTES ---
• CORE NARRATIVE:
* MLO Anchor: [MLO 1]

• SLIDE BODY COVERAGE — KEY MECHANISMS:
Supply Chain Friction:
Plain English: Every delay compounds as orders travel upstream.
"""
    res = we.update_slide(sample_deck, 1, notes)
    assert res["updated"] == 1

    # Verify notes readable via python-pptx
    prs = Presentation(sample_deck)
    read_back = get_slide_notes_text(prs.slides[0])
    assert "⏱ 3.0 min" in read_back
    assert "--- SPEAKER NOTES ---" in read_back
    assert "--- VISUAL DECONSTRUCTION ---" in read_back
    assert "--- LECTURE NOTES ---" in read_back
    assert "Supply Chain Friction" in read_back


def test_write_preserves_an_externally_edited_deck(sample_deck):
    """A deck edited in PowerPoint between two writes must not be reverted.

    Regression: the engine cached its working copy at a fixed /tmp path keyed on
    the deck's *path* and reused it whenever it still existed, so a later write
    republished that stale copy and destroyed edits made in between.
    """
    we.update_slide(sample_deck, 1, _notes())

    _write_deck(sample_deck, ["Introduction", "Core Principles", "USER ADDED SLIDE"])
    assert len(Presentation(sample_deck).slides) == 3

    res = we.update_slide(sample_deck, 2, _notes())
    assert res["updated"] == 1

    titles = [s.shapes.title.text for s in Presentation(sample_deck).slides]
    assert titles == ["Introduction", "Core Principles", "USER ADDED SLIDE"]


def test_publish_refuses_when_the_deck_changes_mid_write(sample_deck):
    """The original must not be replaced if it changed while notes were being written."""
    staging = we.begin_staging(sample_deck)
    try:
        we.inject_notes(staging.deck, [(1, _notes())])
        _write_deck(sample_deck, ["Introduction", "Core Principles", "CHANGED UNDERNEATH"])

        with pytest.raises(RuntimeError, match="changed on disk"):
            we.publish_staged(staging)
    finally:
        we.discard_staging(staging)

    titles = [s.shapes.title.text for s in Presentation(sample_deck).slides]
    assert titles == ["Introduction", "Core Principles", "CHANGED UNDERNEATH"]


def test_backup_refreshed_when_content_changes_despite_older_mtime(sample_deck):
    """Backup freshness must be decided by content, not mtime.

    Copy-based restores (cp -p, rsync -a, Drive, Time Machine, git checkout) all
    preserve mtime, so an mtime test treats a *replaced* deck as already backed up
    and leaves the only backup predating the file it was meant to protect.
    """
    first = we.backup_once(sample_deck)

    _write_deck(sample_deck, ["REPLACEMENT DECK"])
    older = os.path.getmtime(first) - 3600
    os.utime(sample_deck, (older, older))

    second = we.backup_once(sample_deck)
    assert second != first
    assert [s.shapes.title.text for s in Presentation(second).slides] == ["REPLACEMENT DECK"]


def test_backup_reused_while_content_is_unchanged(sample_deck):
    """An unchanged deck must not accumulate a fresh multi-MB backup on every write."""
    assert we.backup_once(sample_deck) == we.backup_once(sample_deck)


def test_empty_notes_are_rejected_and_the_deck_is_untouched(sample_deck):
    """An empty payload used to erase a slide's existing notes and still exit 0."""
    before = open(sample_deck, "rb").read()

    with pytest.raises(ValueError, match="Refusing to write empty notes"):
        we.update_slide(sample_deck, 1, "   \n  ")

    assert open(sample_deck, "rb").read() == before
