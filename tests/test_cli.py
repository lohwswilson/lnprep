"""Tests for lnprep CLI commands using Typer CliRunner."""

import json
import os
import shutil
import tempfile
import pytest
from pptx import Presentation
from typer.testing import CliRunner

from lnprep.main import app

runner = CliRunner()


@pytest.fixture
def test_environment():
    tmp_dir = tempfile.mkdtemp(prefix="lnprep-cli-test-")
    deck_path = os.path.join(tmp_dir, "cli_test_deck.pptx")
    prs = Presentation()
    slide_layout = prs.slide_layouts[1]
    s1 = prs.slides.add_slide(slide_layout)
    s1.shapes.title.text = "Module Introduction"
    s1.shapes.placeholders[1].text = "First core concept"
    prs.save(deck_path)

    yield {"dir": tmp_dir, "deck": deck_path}
    shutil.rmtree(tmp_dir, ignore_errors=True)


def test_cli_version():
    result = runner.invoke(app, ["--version"])
    assert result.exit_code == 0
    assert "lnprep version 1.0.0" in result.stdout


def test_cli_help():
    result = runner.invoke(app, ["--help"])
    assert result.exit_code == 0
    assert "Usage: lnprep" in result.stdout
    assert "init" in result.stdout
    assert "audit" in result.stdout
    assert "verify" in result.stdout
    assert "write" in result.stdout
    assert "cache" in result.stdout


def test_cli_init(test_environment):
    env_dir = test_environment["dir"]
    result = runner.invoke(app, ["init", env_dir, "--class-size", "45", "--json"])
    assert result.exit_code == 0
    data = json.loads(result.stdout)
    assert data["success"] is True
    assert os.path.exists(os.path.join(env_dir, "LECTURE_NOTES_GUIDE.md"))


def test_cli_extract(test_environment):
    deck = test_environment["deck"]
    result = runner.invoke(app, ["extract", deck, "--json"])
    assert result.exit_code == 0
    data = json.loads(result.stdout)
    assert data["total_slides"] == 1
    assert data["slides"][0]["title"] == "Module Introduction"


def test_cli_generate_prompt_only(test_environment):
    deck = test_environment["deck"]
    result = runner.invoke(app, ["generate", deck, "--slide", "1", "--json"])
    assert result.exit_code == 0
    data = json.loads(result.stdout)
    assert "prompt" in data
    assert "Module Introduction" in data["prompt"]


def test_cli_cache_operations(test_environment):
    env_dir = test_environment["dir"]

    # Record
    rec_res = runner.invoke(
        app,
        [
            "cache",
            "record",
            env_dir,
            "Christopher",
            "2016",
            "--title",
            "Logistics & Supply Chain Management",
            "--url",
            "https://example.org/book",
            "--json",
        ],
    )
    assert rec_res.exit_code == 0

    # Lookup
    lookup_res = runner.invoke(
        app,
        ["cache", "lookup", env_dir, "Christopher", "2016", "--json"],
    )
    assert lookup_res.exit_code == 0
    lookup_data = json.loads(lookup_res.stdout)
    assert lookup_data["status"] == "HIT_FRESH"

    # Report
    rep_res = runner.invoke(app, ["cache", "report", env_dir, "--json"])
    assert rep_res.exit_code == 0
    rep_data = json.loads(rep_res.stdout)
    assert rep_data["fresh"] == 1

    # Prune dry-run
    prune_res = runner.invoke(app, ["cache", "prune", env_dir, "--dry-run", "--json"])
    assert prune_res.exit_code == 0


def test_cli_write_and_audit(test_environment):
    deck = test_environment["deck"]
    notes = """--- SPEAKER NOTES ---
⏱ 2.0 min | 🎯 Introduction
BRIDGE: "Welcome to class."
KEY POINT: "Understanding operations is critical for managers."

--- LECTURE NOTES ---
• CORE NARRATIVE:
* MLO Anchor: [MLO 1]

• SLIDE BODY COVERAGE — CONCEPTS:
First core concept:
Plain English: This represents the fundamental basis of our inquiry.
"""
    write_res = runner.invoke(
        app,
        ["write", deck, "1", "--notes", notes, "--allow-unverified", "--json"],
    )
    assert write_res.exit_code == 0
    write_data = json.loads(write_res.stdout)
    assert write_data["success"] is True

    # Audit command
    audit_res = runner.invoke(app, ["audit", deck, "--json"])
    assert audit_res.exit_code == 0
    audit_data = json.loads(audit_res.stdout)
    assert len(audit_data["results"]) == 1


def test_cli_enhance(test_environment):
    deck = test_environment["deck"]
    res = runner.invoke(app, ["enhance", deck, "--slide", "1", "--json"])
    assert res.exit_code == 0
    data = json.loads(res.stdout)
    assert len(data["results"]) == 1


def test_cli_verify(test_environment):
    deck = test_environment["deck"]
    res = runner.invoke(app, ["verify", deck, "--offline", "--json"])
    assert res.exit_code == 0
    data = json.loads(res.stdout)
    assert "counts" in data


def test_cli_cross_check(test_environment):
    env_dir = test_environment["dir"]
    res = runner.invoke(app, ["cross-check", env_dir, "--json"])
    assert res.exit_code == 0
    data = json.loads(res.stdout)
    assert "status" in data


def test_cli_sync(test_environment):
    env_dir = test_environment["dir"]
    res = runner.invoke(app, ["sync", env_dir, "--json"])
    assert res.exit_code == 0
    data = json.loads(res.stdout)
    assert "module_learning_outcomes" in data


def test_cli_brief(test_environment):
    deck = test_environment["deck"]
    notes = '--- SPEAKER NOTES ---\n⏱ 2.0 min | 🎯 Intro\nKEY POINT: "Valid core."\n\n--- LECTURE NOTES ---\n• CORE NARRATIVE:\n* MLO Anchor: [MLO 1]'
    runner.invoke(app, ["write", deck, "1", "--notes", notes, "--allow-unverified"])
    res = runner.invoke(app, ["brief", deck, "--skip-images", "--json"])
    assert res.exit_code == 0
    data = json.loads(res.stdout)
    assert data["stats"]["slides"] >= 1


def test_cli_check(test_environment):
    deck = test_environment["deck"]
    res = runner.invoke(app, ["check", deck, "--json"])
    data = json.loads(res.stdout)
    assert "audit" in data
    assert "references" in data

    # The fixture deck has body text but no notes at all, so the audit half must fail
    # and the gate must report it. Previously audit_failures was always 0 here and the
    # command exited 0 on any deck, whatever state it was in.
    assert data["pass"] is False
    assert data["audit"]["failures"] >= 1
    assert res.exit_code == 1


def test_cli_audit_summary_is_populated(test_environment):
    """Regression: run_audit returned no 'summary', so this table printed all zeros."""
    res = runner.invoke(app, ["audit", test_environment["deck"], "--json"])
    assert res.exit_code == 0
    summary = json.loads(res.stdout)["summary"]
    assert sum(summary.values()) == 1
    assert summary["no_notes"] == 1


def _write_guide_module(tmp_dir):
    session = os.path.join(tmp_dir, "MO9529.Demo", "Session_1")
    os.makedirs(session)
    with open(os.path.join(session, "LECTURE_NOTES_GUIDE.md"), "w") as f:
        f.write(
            "# Guide\n\n## Session Identity\n\n```yaml\nmodule_code: \"MO9529\"\n"
            "duration: \"180 min\"\n```\n\n## Suggested Session Flow\n\n"
            "### Block 1 — Opening (30 min)\n\n## Formatting & Notes Rules\n\n- x\n"
        )
    open(os.path.join(session, "Orphan.pptx"), "w").close()
    return os.path.join(tmp_dir, "MO9529.Demo")


def test_cli_cross_check_reports_failure_and_exits_nonzero():
    """Regression: the command always printed 'status: UNKNOWN' and exited 0, even
    when the guide's time budget and its file references had both failed."""
    tmp_dir = tempfile.mkdtemp(prefix="lnprep-cc-test-")
    try:
        module = _write_guide_module(tmp_dir)
        res = runner.invoke(app, ["cross-check", module, "--json"])
        data = json.loads(res.stdout)

        assert data["status"] == "FAIL"
        assert res.exit_code == 1
        assert data["sessions"]["Session_1"]["budget"]["status"] == "FAIL"
        assert data["sessions"]["Session_1"]["deck"]["status"] == "FAIL"
    finally:
        shutil.rmtree(tmp_dir, ignore_errors=True)


def test_cli_cross_check_rejects_an_unknown_check_name():
    """An unrecognised --check value used to match no branch and silently run nothing."""
    result = runner.invoke(app, ["cross-check", os.getcwd(), "--check", "nonsense"])
    assert result.exit_code == 2
    combined = result.stdout + getattr(result, "stderr", "")
    assert "Unknown --check value" in combined

